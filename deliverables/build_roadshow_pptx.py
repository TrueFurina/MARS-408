# -*- coding: utf-8 -*-
"""Build the 芒得很职 国创赛 5-min roadshow PPT (.pptx) from the WS-B script.

源脚本：deliverables/芒得很职-国创赛路演PPT脚本-5分钟黄金路径.md（11 页，与本脚本 11 页一一对应）
依赖：python-pptx。运行：
    py-server/.venv/Scripts/python.exe deliverables/build_roadshow_pptx.py

产物（两份同源，内容完全一致）：
1. deliverables/芒得很职-国创赛路演PPT-5分钟黄金路径.pptx   —— 物料线留档
2. submission/01_演示PPT/作品演示PPT-最终版.pptx            —— 对外提交件

口径铁律（改动前务必核对 submission/00_提交清单.md）：
- 测试：默认集 1389（1397 collected，addopts deselect 8）→ 1178 passed / 208 skipped / 3 xfailed / 0 failed；CI p0 档 67 条
- 覆盖率 54.96%（门禁 --cov-fail-under=50）
- API：244 operations / 227 paths（app.openapi() 运行时实测）；认证覆盖率 94.67%（231/244）
- 共识分 mean 7.0446 / std 0.2857 —— 必须同时标 stratified 分布 / 96 样本 / seed 20261006 / n=6
- 禁用：451/243 端点、35 路由、1133/915/923/1139/1177/1168、54.08%、6.8413、旧品牌名与旧赛事名
"""
import os

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

_HERE = os.path.dirname(os.path.abspath(__file__))
_ROOT = os.path.dirname(_HERE)

#: 产物路径（两份同源）
OUT_PATHS = (
    os.path.join(_HERE, "芒得很职-国创赛路演PPT-5分钟黄金路径.pptx"),
    os.path.join(_ROOT, "submission", "01_演示PPT", "作品演示PPT-最终版.pptx"),
)

NAVY  = RGBColor(0x0F, 0x2A, 0x43)
TEAL  = RGBColor(0x12, 0x9A, 0x9A)
AMBER = RGBColor(0xE6, 0x7E, 0x22)
LIGHT = RGBColor(0xF4, 0xF7, 0xFA)
AMBBG = RGBColor(0xFD, 0xF3, 0xE9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK  = RGBColor(0x1A, 0x25, 0x33)
GREY  = RGBColor(0x5A, 0x6B, 0x7B)
GREEN = RGBColor(0x27, 0xAE, 0x60)
AMBTX = RGBColor(0xC0, 0x5A, 0x10)
GRY   = RGBColor(0x7F, 0x8C, 0x8D)
PALE  = RGBColor(0xB8, 0xC4, 0xCE)
FONT  = "Microsoft YaHei"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def new_slide(bg):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = bg
    return s


def band(s, color, h=1.25):
    b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Inches(h))
    b.fill.solid(); b.fill.fore_color.rgb = color
    b.line.fill.background()
    b.shadow.inherit = False
    return b


def title(s, text, top=0.18, size=30, color=WHITE):
    tb = s.shapes.add_textbox(Inches(0.6), Inches(top), Inches(12.1), Inches(0.95))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = text
    p.font.size = Pt(size); p.font.bold = True; p.font.color.rgb = color; p.font.name = FONT
    return tb


def body(s, lines, top, left=0.7, width=11.9, height=5.4, size=18, color=DARK, gap=10, colors=None):
    tb = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame; tf.word_wrap = True
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ln
        p.font.size = Pt(size); p.font.color.rgb = (colors[i] if colors else color); p.font.name = FONT
        p.space_after = Pt(gap); p.alignment = PP_ALIGN.LEFT; p.line_spacing = 1.15
    return tb


def footer(s, page):
    tb = s.shapes.add_textbox(Inches(0.6), Inches(7.05), Inches(9), Inches(0.35))
    p = tb.text_frame.paragraphs[0]
    p.text = "芒得很职 · 计算机类学生职业素养对抗实训平台 · 国创赛路演 · 闽江大学"
    p.font.size = Pt(10); p.font.color.rgb = GREY; p.font.name = FONT
    pn = s.shapes.add_textbox(Inches(12.2), Inches(7.05), Inches(0.9), Inches(0.35))
    pp = pn.text_frame.paragraphs[0]; pp.text = str(page)
    pp.font.size = Pt(10); pp.font.color.rgb = GREY; pp.font.name = FONT; pp.alignment = PP_ALIGN.RIGHT


def notes(s, text):
    s.notes_slide.notes_text_frame.text = text


# ---------- Slide 1: Cover ----------
s = new_slide(NAVY)
tb = s.shapes.add_textbox(Inches(1.0), Inches(2.3), Inches(11.3), Inches(1.4))
tf = tb.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; p.text = "芒得很职"
p.font.size = Pt(60); p.font.bold = True; p.font.color.rgb = WHITE; p.font.name = FONT; p.alignment = PP_ALIGN.CENTER
tb2 = s.shapes.add_textbox(Inches(1.0), Inches(3.9), Inches(11.3), Inches(1.0))
p2 = tb2.text_frame.paragraphs[0]; p2.text = '让"软素养"从主观印象，变成可考 · 可溯 · 可练的客观证据链'
p2.font.size = Pt(24); p2.font.color.rgb = TEAL; p2.font.name = FONT; p2.alignment = PP_ALIGN.CENTER
tb3 = s.shapes.add_textbox(Inches(1.0), Inches(6.2), Inches(11.3), Inches(0.8))
p3 = tb3.text_frame.paragraphs[0]
p3.text = "闽江大学  |  新一代多智能体赋能的计算机类学生职业素养对抗实训平台  |  国创赛·高教主赛道·创意组"
p3.font.size = Pt(15); p3.font.color.rgb = PALE; p3.font.name = FONT; p3.alignment = PP_ALIGN.CENTER
notes(s, '面试官问你"抗压能力怎么样"，你只能说"还行"。芒得很职干的事，是把这场对话搬进 AI 多轮对抗，让每一轮交锋都变成一条带原话的证据，结束后给你一份 6 维能力报告。一句话：让软素养从主观印象，变成可考、可溯、可练的客观证据链。')

# ---------- Slide 2: Pain points ----------
s = new_slide(LIGHT); band(s, TEAL)
title(s, "痛点：软素养为什么难练")
body(s, [
    '• 软素养长期靠"老师印象分"，缺结构化、可复现评估',
    '• 真实高压场景（面试 / 需求评审 / 故障通报）学生极少有低成本反复练的机会',
    '• 传统测评只给总分，学生不知道"哪句话扣分、依据是什么"',
    '• 一对一线下模拟成本高，难覆盖一个班级常态化训练',
], top=1.7)
footer(s, 2)
notes(s, '计算机专业学生最怕的往往不是写代码，而是面试、需求谈判、线上故障通报这些高压时刻。这些能力长期靠老师"印象分"评判，不可复现；真实场景又贵又稀缺，练一次的成本极高；就算测了，也只给一个总分，学生根本不知道自己哪句话扣的分。这是所有实训平台的共性难题。')

# ---------- Slide 3: Positioning ----------
s = new_slide(LIGHT); band(s, NAVY)
title(s, "产品定位：多智能体对抗实训")
body(s, [
    '• 不是"刷题 / 背题"工具，而是企业真实场景的 AI 多轮对抗',
    '• 学生 vs AI 面试官多轮交锋，每一轮对话 = 一条可溯源证据',
    '• 自研多智能体技术底座（已工程化落地）：',
    '    - 前端 Vue3 + 后端 FastAPI + 编排 LangGraph',
    '    - 三层记忆 + FrugalRAG + 证据核查 + 质量闸门',
], top=1.7)
footer(s, 3)
notes(s, '芒得很职不是又一个刷题软件。它把企业真实面试、需求评审、故障应急搬进 AI 多轮对抗：学生和 AI 面试官反复交锋，每一轮对话都自动挂上证据编号。底层是我们自研的多智能体技术底座——角色化编排、三层记忆、检索增强、证据核查加质量闸门，是一套工程化落地过的底座。')

# ---------- Slide 4: Core loop ----------
s = new_slide(LIGHT); band(s, TEAL)
title(s, "核心闭环：6–8 轮对抗 → 6 维报告")
body(s, [
    '• 选场景 → 2 问画像 → 6–8 轮对抗 → 6 维 ECD 评估 → 证据校验 → 提升路径 → 报告',
    '• 后端：5 张实训数据表 + 核心链路 API（含越权 403 拦截）已落地',
    '    - 全站 API 规模：244 operations / 227 paths（app.openapi() 运行时实测）',
    '    - 认证覆盖率 94.67%（231/244）',
    '• 前端：四步状态机 + SVG 六维雷达 + 证据链回放',
    '• 状态：上述 ✅ 均已代码级验证',
], top=1.55)
footer(s, 4)
notes(s, '完整闭环已经跑通：选场景、两问画像、6 到 8 轮对抗、按 6 维 ECD 软素养评估、证据一致性校验、最后给出提升路径和报告。后端五张数据表、核心链路接口、包括越权拦截都已在代码级验证；全站 API 规模是 244 个操作、227 条路径，这是运行时 openapi 实测出来的，不是估的；认证覆盖率 94.67%。前端状态机、六维雷达和证据回放也都落地。这是演示里评委能直接看到的主路径。')

# ---------- Slide 5: ⭐ 鲶鱼加压 (highlight) ----------
s = new_slide(AMBBG); band(s, AMBER)
title(s, '⭐ 鲶鱼加压：专治"背模板平顺过关"')
body(s, [
    '• 问题：学生背稿可"平顺过关"，传统测评看不出',
    '• 机制：检测模板化信号（连续 2 轮雷同 / 维度无差异 / 信息密度低 / 过度自信）→ 自动切鲶鱼加压',
    '• 反事实追问，逼出具体数字与真实取舍；加压满 2 轮回到常规',
    '• 报告明确标"鲶鱼加压轮" = 抗压维度最硬证据',
    '• 状态：规则版 ✅ 已落地；MAPPO 化 ⚠️ 合成环境 3-seed 实测超规则基线，策略部署在途',
], top=1.7, color=DARK)
footer(s, 5)
notes(s, '这一页是我们的全场亮点——鲶鱼加压。学生最会说的话就是"背模板平顺过关"，传统测评根本看不出来。芒得很职会实时检测模板化信号：连续两轮雷同、维度没差异、信息密度低、过度自信——一旦命中，系统自动切到鲶鱼加压模式，做反事实追问，逼你说出具体数字和真实取舍。加压满两轮再回到常规。而且报告里会明确标出"哪几轮是鲶鱼加压"，作为抗压维度最硬的证据。这套规则版已经上线，MAPPO 强化学习版本在合成环境三个随机种子实测都超过规则基线。')

# ---------- Slide 6: Evidence moat ----------
s = new_slide(LIGHT); band(s, NAVY)
title(s, "证据护城河：可解释 · 可溯源 · 不可刷分")
body(s, [
    '• 每个 6 维评分强制绑定 turn_id 原话片段，无证据维度标 insufficient，由 quality_gate 拦截 —— 禁止编分',
    '• evidence_check 节点复核"评分—证据"一致性，剔除无依据评分',
    '• 6 维 BARS 行为锚定：表达逻辑 / 抗压应变 / 方案拆解 / 协作沟通 / 技术汇报 / 问题解决',
], top=1.7)
footer(s, 6)
notes(s, '我们的护城河是"证据链"。每一个六维评分都强制绑定对话里的原话片段，编号可回放；拿不出证据的维度直接标"证据不足"并被质量闸门拦截——也就是说，系统禁止凭空编分。还有一个专门的证据核查节点，复核评分和证据是否一致，把无依据的评分剔掉。配合六维行为锚定，芒得很职相对通用对话练习的核心壁垒就是：可解释、可溯源、不可刷分。')

# ---------- Slide 7: Tech credibility ----------
s = new_slide(LIGHT); band(s, TEAL)
title(s, "技术底座可信：工程化 + 防伪证体系")
body(s, [
    '• 芒得很职 多智能体底座已获国家级大创立项背书',
    '• 质量门禁可复现：默认测试集 1389 项 → 1178 passed / 208 skipped / 3 xfailed / 0 failed',
    '    - CI p0 档 67 条；覆盖率 54.96%（门禁 --cov-fail-under=50）',
    '• 共识引擎实测：mean 7.0446 / std 0.2857（stratified 分布 / 96 样本 / seed 20261006 / n=6）',
    '• 同源知识图谱防伪证流水线已于 2026-09-30 并入 main：136 节点 / 18 边',
    '    - 验证报告自带 8 条机检守护（哈希 / 逐行归属 / 锚点 / 复现命令可移植 + 元守护）',
    '• 含义：同一套"写下的结论必须可被机器复验"的工程纪律，同时保障考研线与职业线',
], top=1.5)
footer(s, 7)
notes(s, '我们不是只有 demo。技术底座已经拿到国家级大创立项，说明工程化能力是被认可的。更关键的是工程纪律：默认测试集 1389 项，1178 通过、208 跳过、3 个预期失败、零失败，覆盖率 54.96%，卡在 50 的门槛上；共识引擎实测均分 7.04、标准差 0.29，96 个样本、固定随机种子、6 个智能体，评委可以自己复现。和底座同源的知识图谱防伪证流水线，136 个节点、18 条边，验证报告自带 8 条机器检查守护，连"复现命令能不能在别人机器跑通"都有守护盯着。同一套"写下的结论必须能被机器复验"的纪律，同时保着两条产品线。这就是我们敢说"可信"的底气。')

# ---------- Slide 8: Innovation ----------
s = new_slide(LIGHT); band(s, NAVY)
title(s, "创新点")
body(s, [
    '1. 范式创新：从"资源一次性生成"转向"人机多轮对抗 + 过程性取证"',
    '2. 评估可信：6 维 BARS + 每分挂原话 + 一致性闸门，杜绝随口给分',
    '3. 抗模板机制：鲶鱼加压动态破解"背稿平顺过关"',
    '4. 技术复用：同一底座衍生考研学习与职业训练两条线',
    '5. 数据资产：过程性对话证据天然沉淀为软素养训练语料，越用越准',
], top=1.7)
footer(s, 8)
notes(s, '总结创新点：第一，范式变了——从一次性生成资源，变成多轮对抗加过程取证；第二，评估可信，六维锚定加原话证据加闸门，杜绝随口给分；第三，鲶鱼抗模板；第四，同一底座复用到考研和职业两条线；第五，每一次对话都沉淀成训练语料，越用越准。')

# ---------- Slide 9: Business model ----------
s = new_slide(LIGHT); band(s, TEAL)
title(s, "商业模式（标注规划）")
body(s, [
    '• B2B2C 院校版：班级实训 + 学情看板 SaaS，按班级/席位订阅   🔲 规划',
    '• 实验室/竞赛集训：面试冲刺集训包   🔲 规划',
    '• 增值报告：6 维溯源报告 + 提升路径 PDF 导出   ✅ 功能已具备，商业化待验证',
    '• 合规：仅留学号后 4 位 + 姓名，对话按 user_id 隔离',
], top=1.7, colors=[GRY, GRY, GREEN, DARK])
footer(s, 9)
notes(s, '商业模式上，我们面向高校计算机类专业做院校版 SaaS，也做竞赛集训包；个人增值报告功能已经具备，商业化还在验证。这里要诚实说明：院校订阅和集训包目前是规划，真实运营数据还没有。合规上我们只留学号后四位和姓名，对话内容严格隔离。')

# ---------- Slide 10: Honesty boundary ----------
s = new_slide(LIGHT); band(s, AMBER)
title(s, "诚实边界：已落地 vs 规划")
body(s, ["✅ 已落地："], top=1.55, size=20, color=GREEN, gap=4)
body(s, [
    '• P0–P2 后端闭环、学生端前端、ECD 评估深化、鲶鱼规则版、增值报告功能',
], top=1.95, size=17, gap=8)
body(s, ["⚠️ 在途："], top=2.55, size=20, color=AMBTX, gap=4)
body(s, [
    '• 鲶鱼 MAPPO 部署、教师端（建班 / 学情看板）',
], top=2.95, size=17, gap=8)
body(s, ["🔲 未启动（诚实边界）："], top=3.55, size=20, color=GRY, gap=4)
body(s, [
    '• 46 人真实试点尚未开展 —— 所有"试点成效"类数字一律不存在',
    '• ⚠️ 真实教学增益未经证实：MAPPO / 解析式评审结论来自合成环境，不表述为真实成效',
], top=3.95, size=17, gap=8)
footer(s, 10)
notes(s, '这一页我们主动亮底牌。已经落地的很清楚：后端闭环、前端、评估深化、鲶鱼规则版、报告功能。在途的是 MAPPO 部署和教师端。必须明示：46 人真实试点还没启动，所以今天没有任何"试点成效"数字；强化学习那些结论来自合成环境，我们绝不把它说成真实教学增益。可信，首先是敢把边界说清楚。')

# ---------- Slide 11: Closing ----------
s = new_slide(NAVY)
tb = s.shapes.add_textbox(Inches(1.0), Inches(2.4), Inches(11.3), Inches(1.2))
p = tb.text_frame.paragraphs[0]; p.text = "芒得很职"
p.font.size = Pt(54); p.font.bold = True; p.font.color.rgb = WHITE; p.font.name = FONT; p.alignment = PP_ALIGN.CENTER
tb2 = s.shapes.add_textbox(Inches(1.0), Inches(3.8), Inches(11.3), Inches(0.9))
p2 = tb2.text_frame.paragraphs[0]; p2.text = '把"软素养"变成可考 · 可溯 · 可练的证据'
p2.font.size = Pt(26); p2.font.color.rgb = TEAL; p2.font.name = FONT; p2.alignment = PP_ALIGN.CENTER
tb3 = s.shapes.add_textbox(Inches(1.5), Inches(5.3), Inches(10.3), Inches(1.0))
p3 = tb3.text_frame.paragraphs[0]
p3.text = "期待与各位评委和院校一起，把这件难而正确的事做下去。谢谢！"
p3.font.size = Pt(18); p3.font.color.rgb = PALE; p3.font.name = FONT; p3.alignment = PP_ALIGN.CENTER
notes(s, '芒得很职要做的，是让计算机专业学生的软素养，第一次变成可考、可溯、可练的客观证据。我们不只是做了一个平台，更建立了一套"结论可被机器复验"的工程纪律。期待与各位评委和院校一起，把这件难而正确的事做下去。谢谢！')

for _out in OUT_PATHS:
    os.makedirs(os.path.dirname(_out), exist_ok=True)
    prs.save(_out)
    print("SAVED:", _out)
print("slides:", len(prs.slides._sldIdLst))
