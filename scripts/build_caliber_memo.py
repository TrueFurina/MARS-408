# -*- coding: utf-8 -*-
"""生成《大创交付件口径校验备忘录》docx —— 把诚实口径真值 + 权威来源锁成一份防漂移证据页。
输出：deliverables/闽江申报材料归档-2026-09/大创交付件口径校验备忘录.docx
"""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = "deliverables/闽江申报材料归档-2026-09/大创交付件口径校验备忘录.docx"

doc = Document()
# 基础字体
style = doc.styles["Normal"]
style.font.name = "宋体"
style.font.size = Pt(10.5)

def h(text, size=13):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = True; r.font.size = Pt(size)
    return p

def para(text, size=10.5, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text); r.font.size = Pt(size); r.bold = bold
    return p

def kv_table(rows):
    t = doc.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    for i, txt in enumerate(["项目", "口径真值（可复现）", "权威来源", "备注"]):
        hdr[i].paragraphs[0].add_run(txt).bold = True
    for r in rows:
        c = t.add_row().cells
        for i, txt in enumerate(r):
            c[i].paragraphs[0].add_run(txt)
    return t

# ── 标题 ──
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
tr = title.add_run("大创交付件口径校验备忘录")
tr.bold = True; tr.font.size = Pt(16)
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sr = sub.add_run("NetLearn / MARS-408 · 闽江大学大学生创新创业训练计划 · 防口径漂移证据页")
sr.font.size = Pt(10); sr.italic = True
datep = doc.add_paragraph()
datep.alignment = WD_ALIGN_PARAGRAPH.CENTER
datep.add_run("生成日期：2026-09-17　|　适用范围：附件1 中期检查报告书 / 附件4 研究报告 / 结题草稿").font.size = Pt(9)

# ── 说明 ──
h("一、目的与红线", 12)
para("本备忘录将大创交付件中所有「数值 / 计数 / 通道 / 架构」口径锁定为可复现的权威真值，防止文档间数字漂移与「旧基准乐观值当现行结论」两类问题。")
para("诚信红线：发生型证据（Trace / 成效 / 已训练）不可造假；结构型可标『v1 规则原型』；知识库扩容到 ≥5000 因缺真实教材来源未做（守红线不凑数）。", bold=False)

# ── 核对表 ──
h("二、口径真值核对表", 12)
kv_table([
    ["知识库规模", "2122 个知识点 chunk（E5-base-v2 768 维真实向量，经验证全 L2 范数=1.0、零零向量）",
     "py-server/vectordb_data/netlearn_kb.json（ids 长度=2122）", "磁盘实测值；旧稿写 2083 已订正"],
    ["知识图谱", "86 个知识点 / 82 条先修依赖边（DAG）",
     "teaching_rules 构建；kg_probe.txt 可复现", "旧稿误写 613 节点/609 边 已订正"],
    ["多智能体编排", "11 节点 LangGraph 流水线（13 个 Agent 角色）",
     "py-server/agents/graph.py:99-109", "旧稿写 10 节点 已订正"],
    ["共识机制（NeuralMixer）", "96 样本权威评测：训练权重准确率 80.2% vs 规则加权投票基线 82.3% 持平（未显著超越）；核心收益=修复退化随机映射 consensus≈−0.48→6.46、注意力-质量 Pearson −0.15→0.22；权重 48/48 可加载",
     "py-server/experiments/results/mixer_before_after_2026-09-03.json（12 题×8 次=96 样本）", "30 题 0.7667→0.8333(+8.7%) 仅为早期小样本探索、已被更大样本复现否定，交付件以 96 样本口径为准"],
    ["E5 检索编码", "真实 E5 编码运行（非 BM25-only 降级）",
     "models/e5-base-v2 存在；py-server/experiments/results/trace_frugalrag_search_2026-09-16.json（distance 0.916/0.937/0.674/0.653/0.643）", "FrugalRAG 真实 Trace 重捕于 2026-09-16，证伪『没了』"],
    ["讯飞通道策略", "后端『讯飞星火 generalv3.5 优先 / DeepSeek 兜底』双通道；X2(spark-x) 端点未授权返回 11200，未启用",
     "config.json / .env：XF_ACTIVE_PRESET=generalv3.5", "旧稿写『X2 优先』红线错误 已订正"],
    ["讯飞能力计数", "9 项 AI 能力（图片理解/万搜聚合搜索/智能PPT/数字人视频/文本纠错/公文校对/内容合规/角色模拟/智能简历）",
     "py-server/api/xfyun.py 路由清单", "旧稿列 10/11 项 已订正"],
    ["检索增强相对提升", "Recall@5 +10.7pp、Precision@5 +16.4pp、MRR +9.6pp；token 基本持平（−0.14%）",
     "py-server/experiments/results/benchmark_2026-08-17.json（28 查询 + 30 题×3 次，seed 20260719）", "同源两口径（绝对/相对）均合法"],
])

# ── 校验结论 ──
h("三、机验结论（2026-09-17）", 12)
para("对附件1 / 附件4 / 结题草稿三份 docx 做全文抽取校验：X2 优先 / 613-609 / 10 节点 / 10-11 项 / 2083 / 未标注旧共识 全部 = 0；96 样本口径、86-82、11 节点、generalv3.5+9 项 均已落地。")
para("剩余 0.7667/0.8333 提及均显式标注为『早期 30 题小样本探索 / 已据实校准』，不作头条准确率。", )

# ── 权威来源索引 ──
h("四、权威来源文件索引", 12)
for s in [
    "py-server/vectordb_data/netlearn_kb.json — KB 向量与计数（2122）",
    "py-server/experiments/results/mixer_before_after_2026-09-03.json — NeuralMixer 96 样本评测",
    "py-server/experiments/results/trace_frugalrag_search_2026-09-16.json + TRACE_FrugalRAG_证据说明-2026-09-16.md — FrugalRAG 真实 Trace",
    "py-server/experiments/results/benchmark_2026-08-17.json — 检索增强基准",
    "py-server/agents/graph.py:99-109 — LangGraph 11 节点",
    "kg_probe.txt — 知识图谱 86 节点/82 边",
    "config.json / .env — 讯飞 generalv3.5 运行配置",
]:
    doc.add_paragraph(s, style="List Bullet")

doc.add_paragraph()
foot = doc.add_paragraph()
fr = foot.add_run("本备忘录由代码/产物反查生成，所有数字均可回溯至上述文件字段；如有来源更新，以最新评测文件为准。")
fr.font.size = Pt(9); fr.italic = True

doc.save(OUT)
print("已生成:", OUT)
